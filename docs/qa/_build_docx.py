"""Build the hire-QA Word pack from the markdown sources."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

OUT = Path(__file__).with_name("Store-Website-QA-Pack.docx")

NAVY = RGBColor(0x1A, 0x2A, 0x44)
GOLD = RGBColor(0x8A, 0x6D, 0x2F)
INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x52, 0x60, 0x7D)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HEADER_BG = "1A2A44"
ALT_ROW = "F4F6FA"
LINE = "D6DCEB"


def set_run(run, *, size=11, bold=False, color=INK, italic=False, font="Calibri"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def shade(cell, hex_color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), LINE)
        tcBorders.append(el)
    tcPr.append(tcBorders)


def set_cell_text(cell, text, *, bold=False, color=INK, size=10, fill=None, center=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run(run, size=size, bold=bold, color=color)
    if fill:
        shade(cell, fill)
    set_cell_border(cell)
    cell.vertical_alignment = 1  # center


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18 if level == 1 else 12)
    p.paragraph_format.space_after = Pt(6)
    if level == 1:
        run = p.add_run(text)
        set_run(run, size=18, bold=True, color=NAVY, font="Calibri")
        # bottom rule via border
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "12")
        bottom.set(qn("w:space"), "4")
        bottom.set(qn("w:color"), "C4A35A")
        pBdr.append(bottom)
        pPr.append(pBdr)
    else:
        run = p.add_run(text)
        set_run(run, size=13, bold=True, color=NAVY)
    return p


def add_body(doc, text, *, bold=False, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run(run, size=11, bold=bold, italic=italic)
    return p


def add_rich(doc, parts):
    """parts: list of (text, kwargs)."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    for text, kwargs in parts:
        run = p.add_run(text)
        set_run(run, size=11, **kwargs)
    return p


def add_bullet(doc, text, *, numbered=False, n=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Cm(0.75)
    prefix = f"{n}. " if numbered else "• "
    run = p.add_run(prefix + text)
    set_run(run, size=11)
    return p


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, color=WHITE, size=9, fill=HEADER_BG)
    for r_i, row in enumerate(rows):
        fill = ALT_ROW if r_i % 2 == 0 else "FFFFFF"
        for c_i, val in enumerate(row):
            set_cell_text(table.rows[r_i + 1].cells[c_i], val, size=9, fill=fill)
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return table


def add_blank_table(doc, headers, n_rows, col_widths=None):
    empty = [[""] * len(headers) for _ in range(n_rows)]
    return add_table(doc, headers, empty, col_widths)


def page_break(doc):
    doc.add_page_break()


def cover(doc):
    for _ in range(4):
        doc.add_paragraph()
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = t.add_run("STORE WEBSITE")
    set_run(run, size=12, bold=True, color=GOLD)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.paragraph_format.space_after = Pt(8)
    run = t.add_run("QA Checklist & Testing Instructions")
    set_run(run, size=28, bold=True, color=NAVY, font="Calibri")

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = t.add_run("For hired testers  ·  One pack per client launch")
    set_run(run, size=12, italic=True, color=MUTED)

    doc.add_paragraph()

    meta = [
        ("Client", "________________________________"),
        ("Preview / production URL", "________________________________"),
        ("Leads live?", "Yes / No   Test phone/email: ______________"),
        ("Admin in scope?", "Yes / No"),
        ("Paid landing slugs", "________________________________"),
        ("Cycle date", "________________________________"),
        ("Tester name", "________________________________"),
    ]
    table = doc.add_table(rows=len(meta), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (k, v) in enumerate(meta):
        set_cell_text(table.rows[i].cells[0], k, bold=True, color=WHITE, size=10, fill=HEADER_BG)
        set_cell_text(table.rows[i].cells[1], v, size=10, fill="FFFFFF")
        table.rows[i].cells[0].width = Inches(2.3)
        table.rows[i].cells[1].width = Inches(4.4)

    doc.add_paragraph()
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = note.add_run(
        "Machine checks catch broken markup. QA exists to catch wrong facts, "
        "dead customer paths, and “it works on my laptop” defects."
    )
    set_run(run, size=10, italic=True, color=MUTED)


def section_instructions(doc):
    add_heading(doc, "Part 1 — Testing instructions", 1)

    add_heading(doc, "1. What you are testing", 2)
    add_body(
        doc,
        "This is a store website (hot tubs / spas and similar). Customers browse inventory, "
        "call or text the shop, and submit a short quiz so the business can send pricing.",
    )
    add_table(
        doc,
        ["Route", "Purpose"],
        [
            ["/", "Homepage (hero, lead card, sections)"],
            ["/hot-tubs, /swim-spas, …", "Category listings (only categories this client sells)"],
            ["/hot-tubs/{slug}", "Product detail"],
            ["/inventory", "All in-stock units"],
            ["/find-your-match", "Full-page quiz"],
            ["/visit-us", "Address, hours, map, directions"],
            ["/financing", "Financing terms, or 404 if this client has none"],
            ["/privacy-policy", "Legal"],
            ["/thank-you", "After a successful lead"],
            ["/lp/{slug}", "Paid ads landing (no main nav)"],
            ["/admin/", "Inventory admin — never linked from the public site"],
        ],
        [2.4, 4.3],
    )
    add_rich(
        doc,
        [
            (
                "There is one header, one footer, and one quiz. If the menu works on Home and fails on Visit Us, that is S1.",
                {"bold": True},
            )
        ],
    )

    add_heading(doc, "2. Before you start", 2)
    add_body(doc, "Ask the hiring manager for:")
    items = [
        "Base URL (preview or production). Never invent one.",
        "Client facts sheet: legal name, phone, address, hours, founding year, categories they actually sell, service areas.",
        "Whether leads are live: preview vs production CRM (GoHighLevel). If live, use the test name/email/phone they provide — never your personal number unless they say so.",
        "Admin password only if you are assigned admin testing. Do not request it otherwise.",
        "Paid landing slugs if any (/lp/...).",
        "Devices required (default matrix below if they do not specify).",
    ]
    for i, t in enumerate(items, 1):
        add_bullet(doc, t, numbered=True, n=i)
    add_rich(
        doc,
        [
            (
                "Confirm the site is not showing a red TEMPLATE MODE banner and not using placeholder facts (555 numbers, example.com, Null Island). If it is, stop public testing and report immediately.",
                {"bold": True},
            )
        ],
    )

    add_heading(doc, "3. How to work", 2)
    for t in [
        "Test like a customer first, then like a skeptic (wrong URLs, empty states, double-tap submit).",
        "Complete one area of the checklist, then file bugs, then continue. Do not batch 40 issues in your head.",
        "Every bug needs a URL. “The button is broken” is not a report.",
        "If a defect appears on more than one page, say so (same header/footer/quiz = likely one root cause).",
        "Do not “fix” copy you think is ugly. Report it. Truth of facts (year, phone, prices) is in scope; taste is only in scope if it blocks use (unreadable text, overlapping buttons).",
        "Do not scrape, brute-force admin, or test security beyond: “admin is not linked from the footer” and “wrong password is rejected.”",
    ]:
        add_bullet(doc, t)

    add_heading(doc, "Pass / fail / blocked", 2)
    add_table(
        doc,
        ["Mark", "Meaning"],
        [
            ["P — Pass", "You performed the step and the expected result happened."],
            ["F — Fail", "File a bug. Leave the item unchecked and put the bug ID next to it."],
            ["B — Blocked", "You could not run the step (no admin access, CRM not connected). Say what blocked you."],
        ],
        [1.6, 5.1],
    )

    add_heading(doc, "4. Device and browser matrix", 2)
    add_body(
        doc,
        "Run the smoke pass (section 6) on every required cell. Run the full checklist on at least one desktop and one real phone.",
    )
    add_table(
        doc,
        ["Viewport", "Chrome", "Safari", "Firefox", "Edge"],
        [
            ["Desktop ~1440px", "Required", "macOS if available", "Nice to have", "Windows required"],
            ["Tablet ~768px", "Required", "iPad if available", "—", "—"],
            ["Phone ~390px", "Android Chrome required", "iPhone Safari required", "—", "—"],
        ],
        [1.5, 1.5, 1.6, 1.1, 1.4],
    )
    add_body(
        doc,
        "Facebook in-app browser (iOS) is the highest-risk form traffic. If you can open the URL from a Facebook ad preview or Messenger, do the quiz there once.",
    )
    add_body(doc, "Always test:")
    for t in [
        "Mobile hamburger: tap to open, tap a link, tap overlay/close, confirm the page does not stay locked from scrolling.",
        "Header Call (tel:) and Text (sms:) on a real phone.",
        "Form autofill for name, phone, email on iPhone Safari.",
    ]:
        add_bullet(doc, t)

    add_heading(doc, "5. Severity (use these labels)", 2)
    add_table(
        doc,
        ["Severity", "Meaning", "Examples"],
        [
            [
                "S1 — Blocker",
                "Customer cannot enquire, call, or trust the business",
                "Quiz never submits; thank-you never appears; phones untappable; site down; TEMPLATE MODE on a live URL",
            ],
            [
                "S2 — High",
                "Major path broken or false",
                "Wrong phone/address; hamburger dead on some pages; lead saved but user sees an error; sold unit looks buyable as new; financing rate with no disclaimer",
            ],
            [
                "S3 — Medium",
                "Workaround exists",
                "Broken product image; a nav label is wrong; map pin slightly off",
            ],
            [
                "S4 — Low",
                "Polish",
                "Alignment, hover, minor copy typo that does not change meaning",
            ],
        ],
        [1.5, 2.2, 3.0],
    )
    add_body(doc, "If you are unsure, pick the higher severity.")

    add_heading(doc, "6. Smoke pass (15–20 minutes)", 2)
    add_body(doc, "Do this at the start of every build/URL you are given.")
    smoke = [
        "Home loads; logo visible; no TEMPLATE banner.",
        "Click every header link; each goes where the label says.",
        "Repeat every header link from the mobile menu.",
        "Tap/click the header phone — it is a tel: link, not selected text.",
        "Primary CTA (for example “Shop Inventory”) reaches /inventory.",
        "Open one category, one product, /visit-us, /find-your-match, /privacy-policy.",
        "Submit the quiz with the test contact (if allowed) and land on /thank-you.",
        "Footer: no Admin link; social icons open real profiles; map iframe shows the shop; privacy link works.",
        "Visit a garbage URL (e.g. /this-page-does-not-exist) — branded 404, useful next steps, same phone number.",
    ]
    for i, t in enumerate(smoke, 1):
        add_bullet(doc, t, numbered=True, n=i)
    add_rich(
        doc,
        [("If smoke fails, do not spend hours on visual polish. File S1/S2 and wait.", {"bold": True})],
    )

    add_heading(doc, "7. Lead / quiz testing (critical path)", 2)
    add_body(
        doc,
        "The quiz is two steps: pick a category (or “Just browsing”), then name / phone / email.",
    )
    add_heading(doc, "Happy path", 2)
    happy = [
        "Prefer a campaign URL when asked to test ads: https://{site}/find-your-match?utm_source=facebook&utm_campaign=qa&fbclid=QA-TEST-001",
        "Choose a category this client sells. Confirm the form shows “Looking at: {that label}” and that Change returns to step 1.",
        "Fill name, phone (10+ digits), email. Optional message is optional.",
        "Submit once. Button should show Sending… and disable (no double posts).",
        "You land on /thank-you?ref=...",
        "If CRM testing is in scope: contact appears in GoHighLevel with tags such as website-lead, category tag, and UTM/source fields. Ask the operator to confirm — testers often cannot see GHL.",
        "Thank-you still offers a working Call link.",
    ]
    for i, t in enumerate(happy, 1):
        add_bullet(doc, t, numbered=True, n=i)

    add_heading(doc, "Negative and edge", 2)
    add_table(
        doc,
        ["Test", "Expected"],
        [
            ["Submit empty required fields", "Stays on form; does not hit thank-you"],
            ["Email without @", "Error or browser validation; no thank-you"],
            ["Phone with fewer than 10 digits", "Rejected with a clear error and a call link"],
            ["Double-click submit", "One lead, or second request is harmless — not two identical CRM contacts from one click"],
            ["Just browsing", "Lead still captured; not tagged as a category they do not sell"],
            ["Quiz on homepage and /find-your-match", "Both submit; source_page should differ (operator/CRM)"],
            ["Quiz on a product page", "Lead associated with that product if CRM is in scope"],
            ["Hidden “Company” field (honeypot)", "Do not fill unless asked. Humans should never see it."],
        ],
        [2.6, 4.1],
    )
    add_body(doc, "Do not paste real customer PII into tickets. Mask phone/email in screenshots.")

    add_heading(doc, "8. Inventory and products", 2)
    add_body(
        doc,
        "Public catalog should only show statuses available, pending, and sold. Draft and hidden must not appear on /inventory or category pages.",
    )
    add_table(
        doc,
        ["Check", "Expected"],
        [
            ["Listing vs detail price", "Same cash and monthly strings (or both “Ask for current pricing”)"],
            ["Pending / sold", "Clearly labelled; look muted, not like a fresh “buy now”"],
            ["Quantity 1 available", "“Last one” (or equivalent badge) consistent on card and detail"],
            ["Images", "Main + gallery load; no broken icons; alt text is the product or business name, not empty"],
            ["Wrong category in URL", "404 “not on the floor”"],
            ["Category this client does not sell", "404, and no nav or 404-page button advertising it"],
        ],
        [2.4, 4.3],
    )
    add_body(
        doc,
        "If the database is not connected, inventory should show a configured empty/error message, not a blank crash or a 500 page.",
    )

    add_heading(doc, "9. Admin (only if assigned)", 2)
    add_body(doc, "/admin/ is typed in the address bar. It must not appear in the footer, header, or sitemap.")
    admin = [
        "Logged-out view: login form only — no product table.",
        "Wrong password: error, still logged out.",
        "Right password: product list loads.",
        "Create a product as draft — confirm it is absent from the public site.",
        "Set available — it appears in the right category and /inventory.",
        "Image: upload under 6 MB JPG/PNG/WEBP/GIF; public page shows it.",
        "Sign out; refresh /admin/ — logged out again.",
    ]
    for i, t in enumerate(admin, 1):
        add_bullet(doc, t, numbered=True, n=i)
    add_body(doc, "Do not delete live inventory unless given a disposable test product.")

    add_heading(doc, "10. Paid landing pages (/lp/...)", 2)
    add_body(doc, "If slugs are provided:")
    for t in [
        "Page has an advertorial / advertising disclosure at the top.",
        "No full site nav (by design — ads should not leak to /inventory before convert).",
        "Form still submits to thank-you.",
        "View source / inspector: noindex.",
        "Phone still works.",
        "Footnotes/disclosures at the bottom are readable.",
    ]:
        add_bullet(doc, t)

    add_heading(doc, "11. Accessibility and layout (minimum)", 2)
    add_body(doc, "Not a full WCAG audit unless contracted. Still fail:")
    for t in [
        "Text over busy photos with contrast so low you cannot read the headline.",
        "Buttons or links smaller than ~44px on mobile, or overlapping the hamburger.",
        "Focus outline completely missing when tabbing the quiz (desktop).",
        "Horizontal scroll on a 390px-wide phone on Home, Inventory, Find Your Match, Visit Us.",
        "Map iframe with no title, or social links with no accessible name.",
    ]:
        add_bullet(doc, t)

    add_heading(doc, "12. How to file a bug", 2)
    add_body(doc, "Use one ticket per defect. Title format:")
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    run = p.add_run("[S2] [mobile Safari] Quiz submit stays on Sending… — /find-your-match")
    set_run(run, size=11, italic=True, color=NAVY)
    add_body(doc, "Body must include:")
    must = [
        "URL (full, including query string if relevant).",
        "Device / OS / browser / viewport.",
        "Steps numbered from a cold load (or “cookies cleared”).",
        "Expected vs actual.",
        "Screenshot or short video (required for visual and layout bugs).",
        "Console errors (F12 → Console), if any.",
        "Severity from section 5.",
        "Happens on which pages (one page vs all pages with header).",
    ]
    for i, t in enumerate(must, 1):
        add_bullet(doc, t, numbered=True, n=i)
    add_body(doc, "A blank bug report form is in Part 3 of this document.")

    add_heading(doc, "13. What “done” means for a pass", 2)
    add_body(doc, "A testing cycle is complete when:")
    for t in [
        "Smoke pass is green on the required matrix.",
        "The full checklist in Part 2 is filled (pass / fail+ticket / blocked).",
        "All S1 and S2 bugs are filed; S3/S4 filed or listed as deferred with the manager’s OK.",
        "At least one successful test lead is confirmed by the operator (or explicitly marked blocked).",
        "You listed content-truth issues separately (wrong year, leftover client name, category they do not sell) — these ship even when the site “works.”",
    ]:
        add_bullet(doc, t)

    add_heading(doc, "14. Out of scope unless asked", 2)
    for t in [
        "Load / penetration testing.",
        "Changing production inventory or CRM automations.",
        "Deploying, merging code, or running wrangler / gate (operators).",
        "Pixel / Zaraz configuration in Cloudflare (operators). If leads duplicate in Meta Ads Manager, report the symptom; do not reconfigure tags.",
    ]:
        add_bullet(doc, t)


def section_checklist(doc):
    add_heading(doc, "Part 2 — QA checklist", 1)
    add_body(
        doc,
        "Mark P (pass), F (fail — write the bug ID), or B (blocked). Do not leave items blank. "
        "Copy this document per cycle, or tick in your tracker.",
    )

    blocks = [
        (
            "A. Launch / truth (human-only — automation will not catch these)",
            [
                ["A1", "No TEMPLATE MODE banner on this URL"],
                ["A2", "Business name matches the client (header, footer, titles, quiz consent)"],
                ["A3", "Phone is this client’s number everywhere (header, footer, visit us, quiz errors, thank-you, inventory empty state)"],
                ["A4", "Phone is always a tappable tel: link, never plain text"],
                ["A5", "Address and city/region match the client; map pin is the shop"],
                ["A6", "Hours match the client (Visit Us + footer + Find Your Match sidebar)"],
                ["A7", "Founded / “since {year}” and copyright year are consistent and true"],
                ["A8", "Categories in nav are only things this client sells (no leftover saunas, chairs, etc.)"],
                ["A9", "Homepage / section copy does not mention products or towns from another client"],
                ["A10", "Nav labels match what this client calls things (one label per destination)"],
                ["A11", "Service-area chips / towns are places they actually serve"],
                ["A12", "Logos: header (light bar) readable; drawer + footer knockout/light logo readable on dark"],
                ["A13", "Favicon loads"],
                ["A14", "Nothing internal visible: no /admin link, no TODO, no CONFIRM, no staging passwords"],
            ],
        ),
        (
            "B. Header, footer, global chrome",
            [
                ["B1", "Desktop: every nav item + primary CTA"],
                ["B2", "Mobile: hamburger opens/closes; links work on Home"],
                ["B3", "Mobile hamburger works on Inventory, a category page, a product page, Find Your Match, Visit Us, Privacy"],
                ["B4", "Active nav state looks correct on the current page"],
                ["B5", "Header Text uses sms: and opens the messenger on a phone"],
                ["B6", "Header CTA label and destination match (e.g. Shop Inventory → /inventory)"],
                ["B7", "Footer Quick links match header destinations (same labels)"],
                ["B8", "Footer legal includes Privacy Policy and it works"],
                ["B9", "Social icons: only real networks; each opens the correct profile"],
                ["B10", "Footer map iframe loads; Get directions (Visit Us) opens maps to the same address"],
                ["B11", "No horizontal overflow at 390px and 1440px on Home"],
            ],
        ),
        (
            "C. Homepage",
            [
                ["C1", "Hero headline/subhead readable on mobile and desktop"],
                ["C2", "Hero / announcement links go to the stated pages"],
                ["C3", "Lead card quiz: category step → form → submit (if leads allowed)"],
                ["C4", "Images (hero, showroom, product shots) load, not stretched or 1px"],
                ["C5", "Buttons have enough tap space; gold/primary CTAs visible"],
                ["C6", "Page does not jump wildly while images load"],
            ],
        ),
        (
            "D. Inventory and categories",
            [
                ["D1", "/inventory lists units or a clear empty/unconfigured message (never a crash)"],
                ["D2", "Category jump links (if more than one category) match counts roughly"],
                ["D3", "Each enabled category URL (e.g. /hot-tubs) lists only that category"],
                ["D4", "Product cards: name, image, price/monthly or “Ask for current pricing”"],
                ["D5", "Card click opens the matching detail URL"],
                ["D6", "Draft products not on public listing"],
                ["D7", "Hidden products not on public listing"],
                ["D8", "Sold / pending still visible but clearly labelled and muted"],
                ["D9", "Disabled category URL 404s and is not in nav or 404 shortcuts"],
            ],
        ),
        (
            "E. Product detail",
            [
                ["E1", "Title, description, badges match the listing card"],
                ["E2", "Cash / monthly / ask-for-price identical to the card"],
                ["E3", "Gallery thumbs switch the main image (if present)"],
                ["E4", "Breadcrumbs: Home → category → product; all work"],
                ["E5", "Quiz on the page submits; thank-you works"],
                ["E6", "Related products (if any) stay in the same category and load"],
                ["E7", "Fake slug /hot-tubs/not-a-real-unit → branded 404"],
                ["E8", "Product in the wrong category segment → 404"],
            ],
        ),
        (
            "F. Find Your Match, financing, visit, legal, thank-you",
            [
                ["F1", "/find-your-match: quiz hydrates (options clickable, not a dead heading)"],
                ["F2", "Quiz options = enabled categories + Just browsing — nothing extra"],
                ["F3", "Sidebar hours + call link match the rest of the site"],
                ["F4", "Browse-instead links match enabled categories"],
                ["F5", "/visit-us: address, hours, directions, phone, email (if published), quiz"],
                ["F6", "/financing: if client has financing, terms + visible disclaimer + quiz; apply URL (if any) opens lender"],
                ["F7", "/financing: if client has no financing, page is 404 and not in nav"],
                ["F8", "/privacy-policy: readable on mobile; contact phone/email match"],
                ["F9", "/thank-you: loads with or without ?ref=; call CTA works"],
                ["F10", "Direct visit to /thank-you without submitting is acceptable (no crash)"],
            ],
        ),
        (
            "G. Lead form quality",
            [
                ["G1", "Labels present for name, phone, email"],
                ["G2", "iPhone Safari autofill fills name / tel / email"],
                ["G3", "Invalid email blocked"],
                ["G4", "Short/invalid phone blocked with error + call link"],
                ["G5", "Submit disabled while Sending…"],
                ["G6", "Network failure (offline) shows error, not a silent thank-you"],
                ["G7", "Consent line names this business"],
                ["G8", "After success, campaign URL test lead confirmed in CRM or marked B"],
                ["G9", "Second submit in same browser updates the same enquiry (no obvious duplicate flood) — operator may confirm"],
            ],
        ),
        (
            "H. 404 and SEO hygiene (spot check)",
            [
                ["H1", "Random path: branded 404, working next-step buttons"],
                ["H2", "/admin/ not listed in the public sitemap"],
                ["H3", "/lp/ URLs (if any) not in the public sitemap"],
                ["H4", "Homepage title and meta description are client-specific, not “Template”"],
                ["H5", "Open Graph image (if any) loads when sharing a product URL"],
            ],
        ),
        (
            "I. Paid landing (/lp/{slug}) — skip if none",
            [
                ["I1", "Advertising / advertorial label visible at top"],
                ["I2", "No main site nav / footer sitemap"],
                ["I3", "Logo present; phone or form can convert"],
                ["I4", "Lead submit → thank-you"],
                ["I5", "Disclosures/footnotes readable"],
                ["I6", "noindex in document head"],
            ],
        ),
        (
            "J. Admin — skip if not assigned",
            [
                ["J1", "URL not linked from public pages"],
                ["J2", "Logged out: no product data in HTML"],
                ["J3", "Bad password rejected"],
                ["J4", "Good password: list loads"],
                ["J5", "Create draft → absent on public site"],
                ["J6", "Publish available → appears in correct category"],
                ["J7", "Edit name/price → listing and detail agree"],
                ["J8", "Image upload appears on the public product page"],
                ["J9", "Logout works"],
            ],
        ),
        (
            "K. Cross-cutting visual / a11y smoke",
            [
                ["K1", "No overlapping header vs hero on mobile"],
                ["K2", "Quiz options readable and tappable (2-col mobile / 3-col desktop)"],
                ["K3", "Keyboard: tab through quiz fields on desktop"],
                ["K4", "No console errors on Home, Inventory, Find Your Match"],
                ["K5", "Print not required; skip unless asked"],
            ],
        ),
    ]

    for title, rows in blocks:
        add_heading(doc, title, 2)
        add_table(
            doc,
            ["#", "Check", "P / F / B", "Bug ID"],
            [[a, b, "", ""] for a, b in rows],
            [0.7, 4.5, 1.0, 0.9],
        )

    add_heading(doc, "Sign-off", 2)
    add_table(
        doc,
        ["Item", "Answer"],
        [
            ["Smoke pass completed on required devices", "Yes / No"],
            ["S1 bugs still open", ""],
            ["S2 bugs still open", ""],
            ["S3 / S4 bugs still open", ""],
            ["Operator confirmed test lead in CRM", "Yes / No / N/A"],
            ["Tester signature", ""],
            ["Date", ""],
        ],
        [3.4, 3.3],
    )


def section_bug(doc):
    add_heading(doc, "Part 3 — Bug report template", 1)
    add_body(doc, "Copy one block per defect. One ticket per defect.")

    add_table(
        doc,
        ["Field", "Fill in"],
        [
            ["Title", "[S#] [device browser] Short symptom — /path"],
            ["Severity", "S1 / S2 / S3 / S4"],
            ["URL", ""],
            ["Environment", "preview / production"],
            ["Client", ""],
            ["Device / OS / browser / width", ""],
            ["Expected", ""],
            ["Actual", ""],
            ["Screenshots / video", ""],
            ["Console / network", ""],
            ["Pages affected", "this page only / all pages with header / all quiz instances"],
            ["Workaround (if any)", ""],
        ],
        [2.2, 4.5],
    )
    add_heading(doc, "Steps to reproduce", 2)
    add_blank_table(doc, ["#", "Action"], 8, [0.6, 6.1])

    add_heading(doc, "Extra copies", 2)
    add_body(doc, "Blank forms for additional bugs:")
    for n in range(1, 4):
        add_heading(doc, f"Bug {n}", 2)
        add_table(
            doc,
            ["Field", "Fill in"],
            [
                ["Title", ""],
                ["Severity", "S1 / S2 / S3 / S4"],
                ["URL", ""],
                ["Device / browser", ""],
                ["Expected", ""],
                ["Actual", ""],
                ["Bug ID (tracker)", ""],
            ],
            [2.2, 4.5],
        )
        add_blank_table(doc, ["#", "Steps"], 5, [0.6, 6.1])


def set_header_footer(doc):
    section = doc.sections[0]
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = hp.add_run("Store website  ·  QA pack for hired testers")
    set_run(run, size=9, color=MUTED)

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run("Confidential to the hiring team  ·  Do not publish admin passwords in this file  ·  Page ")
    set_run(run, size=8, color=MUTED)
    # PAGE field
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run2 = fp.add_run()
    run2._r.append(fld1)
    run2._r.append(instr)
    run2._r.append(fld2)
    set_run(run2, size=8, color=MUTED)


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.page_width = Cm(21.59)
    section.page_height = Cm(27.94)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.font.color.rgb = INK

    set_header_footer(doc)
    cover(doc)
    page_break(doc)
    section_instructions(doc)
    page_break(doc)
    section_checklist(doc)
    page_break(doc)
    section_bug(doc)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
